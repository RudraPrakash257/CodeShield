"""DOCX text extraction."""
from docx import Document
from pathlib import Path


def extract_text_from_docx(file_path: str) -> dict:
    """
    Extract text from DOCX file.

    Returns:
        dict with filename, text, pages, is_handwritten, ocr_used
    """
    try:
        doc = Document(file_path)

        # Extract paragraphs
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Extract tables
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    table_texts.append(row_text)

        all_text = "\n".join(paragraphs + table_texts)

        return {
            "filename": Path(file_path).name,
            "text": all_text,
            "pages": max(1, len(paragraphs) // 10),  # Rough estimate
            "is_handwritten": False,
            "ocr_used": False
        }

    except Exception as e:
        raise Exception(f"Failed to extract DOCX: {str(e)}")
